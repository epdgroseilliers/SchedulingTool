"""ui/scheduling/tag.py — the Tag tab of the link popup, through the
Scheduling View page.

The board is a custom component, so clicking a link reaches Python the way
any board event does: `session_state["mv_board"] = {...}` — see
tests/ui/test_scheduling_view.py's own docstring. Every generated file goes
to tmp_path via `redirect_tags`; nothing here touches Y:\\West.
"""

import copy
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

import data.tagfiles.bilateral as tagfile
from dialog_helpers import let_go_of_closed_widgets
from domain.tags import MAG_PSE, day_folder_name
from ui.scheduling.state import BIDFILE_DIALOG, LINK_DIALOG, PENDING_TAG
from ui.scheduling.tag import seed_is_stale, table_height

PAGE_PATH = str(Path(__file__).resolve().parents[2] / "pages" / "1_Scheduling_View.py")
FLOW = date.today() + timedelta(days=1)


#: What the desk's own mapping table answers for the counterparties these
#: tests trade. Stubbed rather than read live so the fast tier keeps working
#: with no connection to the mapping server; the live wiring is covered by
#: tests/data/test_markets.py's db-marked tests.
PSE_CODES = {"AZPS": "APS01", "BPAT": "BPAP01", "EPE": "EPEC01"}


@pytest.fixture
def no_bilateral_db(monkeypatch):
    monkeypatch.setattr(
        "ui.scheduling.state.load_trades_for_flow_date", lambda flow_date: ([], None)
    )
    monkeypatch.setattr(
        "ui.scheduling.state.pse_for_market", lambda market: PSE_CODES.get(market)
    )
    monkeypatch.setattr(
        "ui.scheduling.tag.known_pse_codes", lambda: sorted(PSE_CODES.values()) + [MAG_PSE]
    )
    monkeypatch.setattr("ui.scheduling.tag.reservation", fake_reservation)


@pytest.fixture
def redirect_tags(tmp_path, monkeypatch):
    """The tag root and the trading session the folder is named for.

    Not TEST_MODE: that is off in production now, so leaving it on here
    would exercise a path the desk no longer takes. The session is stubbed
    because the folder is named for it and the live calendar is not this
    test's business.
    """
    monkeypatch.setattr(tagfile, "TAG_FOLDER", tmp_path)
    monkeypatch.setattr(tagfile, "sessions_near", lambda day: {FLOW: [FLOW]})
    return tmp_path


def hourly_trade(direction, counterparty, location, hours, mw, flow_date=FLOW):
    return {
        "direction": direction, "counterparty": counterparty, "location": location,
        "index": None, "price": 73.0, "is_monthly": False,
        "schedule": [(flow_date, h, float(mw)) for h in hours],
    }


BUY_LEG = hourly_trade("Buy", "AZPS", "PALOVERDE500", range(7, 23), 100)
SELL_LEG = hourly_trade("Sell", "BPAT", "MIDC", range(7, 23), 60)


def _in_dialog(at):
    """Tell the page its modal is still open — see test_bidfile.py."""
    at.session_state[BIDFILE_DIALOG] = False
    at.session_state[LINK_DIALOG] = False
    return at


def _run(trades=()):
    at = AppTest.from_file(PAGE_PATH, default_timeout=120)
    at.session_state["trades"] = list(trades)
    return at.run()


def _emit(at, event):
    event = dict(event)
    event["seq"] = at.session_state["mv_last_seq"] + 1
    at.session_state["mv_board"] = event
    return at.run()


def _click(at, label):
    [b for b in _in_dialog(at).button if b.label == label][0].click().run()
    return let_go_of_closed_widgets(at)


def _tagging(trades=(BUY_LEG, SELL_LEG), frm="session:0", to="session:1", market=None):
    """A link, opened again by clicking it — which is where tagging starts."""
    at = _run(trades)
    request = {"type": "link_request", "from": frm}
    request["market" if market else "to"] = market or to
    at = _click(_emit(at, request), "Create link")
    at = _emit(_in_dialog(at), {"type": "link_click", "link_id": "L1"})
    return _in_dialog(at).run()


def _drawing(trades=(BUY_LEG, SELL_LEG)):
    """A link dragged but not yet created — the popup open on it."""
    at = _emit(_run(trades), {
        "type": "link_request", "from": "session:0", "to": "session:1"
    })
    return _in_dialog(at)


def _day_folder(root):
    """Where a tag for this flow date lands: the session's own folder,
    which the redirect_tags fixture stubs to this one day."""
    return root / day_folder_name([FLOW])


def _tag(at, link_id="L1"):
    return at.session_state["mv_tags"][link_id]


def _fill(at, link_id="L1", gca="GWA", lca="BPAT"):
    at.text_input(f"mv_tag_{link_id}_Source_gca").set_value(gca)
    at.text_input(f"mv_tag_{link_id}_Sink_lca").set_value(lca)
    return _in_dialog(at).run()


#: The one OASIS reservation these tests know about — the real one from the
#: desk's September ABEX-SWPW tag. Stubbed so the fast tier needs no
#: connection; the live lookup is covered by tests/data/test_reservations.py.
OASIS = {
    # Still good on the flow date these tests use.
    "109267045": {
        "aref": "109267045",
        "path": "WS/NWMT/NWMT-WAUW/MATL.NWMT-CROSSOVER/",
        "provider": "NWMT", "por": "MATL.NWMT", "pod": "CROSSOVER",
        "status": "CONFIRMED", "ts_class": "FIRM", "granted_mw": 40.0,
        "first_date": FLOW - timedelta(days=2), "last_date": FLOW + timedelta(days=2),
    },
    # Long expired — a number off an old tag, which is the ordinary case.
    "110616287": {
        "aref": "110616287",
        "path": "W/TEPC/TEPC-AZPS/TEPC.SYS-SAGUARO500/",
        "provider": "TEPC", "por": "TEPC.SYS", "pod": "SAGUARO500",
        "status": "CONFIRMED", "ts_class": "FIRM", "granted_mw": 25.0,
        "first_date": FLOW - timedelta(days=90), "last_date": FLOW - timedelta(days=60),
    },
}


def fake_reservation(aref):
    return OASIS.get((aref or "").strip())


def _type_reservation(at, *arefs, link_id="L1"):
    """Add a transmission row per reference, the way the editor reports one."""
    key = _transmission_key(at, link_id)
    at.session_state[key] = {
        "edited_rows": {}, "deleted_rows": [],
        "added_rows": [{"reservation": a} for a in arefs],
    }
    return _in_dialog(at).run()


def _transmission_key(at, link_id="L1"):
    """The editor's key, rebuilt the way ui.scheduling.tag builds it — it
    carries a version so a filled path can be shown without deleting the
    state of a widget still on screen."""
    ver_key = f"mv_tag_{link_id}_transmissions_ver"
    ver = at.session_state[ver_key] if ver_key in at.session_state else 0
    return f"mv_tag_{link_id}_transmissions" + (f"_v{ver}" if ver else "")


def _transmissions(at, link_id="L1"):
    return at.session_state["mv_tags"][link_id]["transmissions"]


def _numbers(at, link_id="L1"):
    """The "# trans" cells that actually got filled in."""
    return [
        row["reservation"] for row in _transmissions(at, link_id) if row["reservation"]
    ]


#: The opening string these two counterparties imply, in their PSE codes.
OPENING_PATH = "??-APS01-MAG001-BPAP01-??"

#: One round later, as a scheduler would paste it back — the desk's own
#: wording, from the TODO.
CHAT_PATH = (
    "HE1-6; Mead MAG (G @ SWPW) (WALC tx PPK>MEAD, 7f 110704717) - MAG - "
    "EEMU - TNSK - MGM sink"
)


def _path_entry_key(at, link_id="L1"):
    """The paste box's key, rebuilt the way ui.scheduling.tag builds it. It
    carries a version because the box empties itself after each round."""
    ver_key = f"mv_tag_{link_id}_path_entry_ver"
    ver = at.session_state[ver_key] if ver_key in at.session_state else 0
    return f"mv_tag_{link_id}_path_entry" + (f"_v{ver}" if ver else "")


def _paste_path(at, text, link_id="L1"):
    at.text_input(_path_entry_key(at, link_id)).set_value(text)
    return _in_dialog(at).run()


def _shown_path(at):
    """What the popup is displaying as the latest path — the copyable block,
    not the box."""
    return [block.value for block in at.code]


def _tag_expander(at):
    return [block for block in at.expander if block.label == "Tag details"][0]


#: One assembled historical tag (the shape data.tags_history.assemble()
#: produces), sharing both of _tagging()'s counterparties (APS01, BPAP01 —
#: AZPS and BPAT's codes under PSE_CODES above) so "Lookup old tags" finds
#: it. The reservation number is the fixture OASIS already knows, so
#: applying this route exercises the existing OASIS fill too.
HISTORICAL_TAG = {
    "tag_index": 1,
    "gca": "GWA",
    "lca": "SWPW",
    "flow_date": date(2026, 9, 1),
    "market_path": [
        {"pse": "APS01", "product": "G-F", "contracts": ""},
        {"pse": "MAG001", "product": "", "contracts": ""},
        {"pse": "BPAP01", "product": "L", "contracts": ""},
    ],
    "physical_path": [
        {"type": "G", "tp": "", "por_name": "RIMROCK", "por_ca": "GWA",
         "pod_name": "", "pod_ca": "", "reservations": []},
        # One of each kind a reservation number can be: still good, long
        # expired, not a reference at all, and one of someone else's.
        {"type": "T", "tp": "NWMT", "por_name": "MATL.NWMT", "por_ca": "NWMT",
         "pod_name": "CROSSOVER", "pod_ca": "NWMT",
         "reservations": [{"contract": "109267045", "product": "7-F"}]},
        {"type": "T", "tp": "TEPC", "por_name": "TEPC.SYS", "por_ca": "TEPC",
         "pod_name": "SAGUARO500", "pod_ca": "TEPC",
         "reservations": [{"contract": "110616287", "product": "7-F"}]},
        {"type": "T", "tp": "EPE", "por_name": "EPE.SYS", "por_ca": "EPE",
         "pod_name": "PVPV", "pod_ca": "EPE",
         "reservations": [{"contract": "EPE5PVPV", "product": "7-F"}]},
        {"type": "T", "tp": "WALC", "por_name": "PPK", "por_ca": "WALC",
         "pod_name": "MEAD", "pod_ca": "WALC",
         "reservations": [{"contract": "999999999", "product": "7-F"}]},
        {"type": "L", "tp": "", "por_name": "", "por_ca": "",
         "pod_name": "SWPW_HUB", "pod_ca": "SWPW", "reservations": []},
    ],
}

#: A route between counterparties this link never touches — shares no PSE
#: with AZPS/BPAT — so it must never surface regardless of anything else
#: about it.
UNRELATED_TAG = {
    "tag_index": 2,
    "gca": "OTHER",
    "lca": "OTHER",
    "flow_date": date(2026, 9, 1),
    "market_path": [
        {"pse": "ABEX", "product": "G-F", "contracts": ""},
        {"pse": "MAG001", "product": "", "contracts": ""},
        {"pse": "EEMU24", "product": "L", "contracts": ""},
    ],
    "physical_path": [],
}


def _stub_history(monkeypatch, records=(), error=None, capture=None):
    def fake(start, stop):
        if capture is not None:
            capture["start"], capture["stop"] = start, stop
        return list(records), error

    monkeypatch.setattr("ui.scheduling.tag.load_west_tags", fake)


def _use_route(at):
    return _click(at, "Use this route")


pytestmark = pytest.mark.usefixtures("no_bilateral_db")


class TestWhereTaggingStarts:
    def test_clicking_a_link_opens_its_tag_alongside_its_schedule(self):
        at = _tagging()
        assert not at.exception, [e.value for e in at.exception]
        assert "L1" in at.session_state["mv_tags"]

    def test_a_link_being_drawn_has_its_tag_fields_too(self):
        # One popup, so the whole job is on screen before the link exists —
        # the tag is filed under PENDING_TAG until Create link gives it an
        # id to be filed under.
        at = _drawing()
        assert list(at.session_state["mv_tags"]) == [PENDING_TAG]
        assert at.text_input("mv_tag_pending_Source_gca") is not None

    def test_what_was_typed_before_the_link_existed_moves_onto_it(self):
        at = _drawing()
        at.text_input("mv_tag_pending_Source_gca").set_value("GWA")
        at = _in_dialog(at).run()
        at = _click(at, "Create link")
        assert list(at.session_state["mv_tags"]) == ["L1"]
        assert _tag(at)["source"]["gca"] == "GWA"

    def test_abandoning_a_link_takes_its_tag_with_it(self):
        # Or the next link drawn would open holding someone else's path.
        at = _drawing()
        at.text_input("mv_tag_pending_Source_gca").set_value("GWA")
        at = _click(_in_dialog(at).run(), "Cancel")
        assert at.session_state["mv_tags"] == {}

    def test_the_tag_goes_when_the_link_does(self):
        at = _click(_tagging(), "Delete")
        assert at.session_state["mv_tags"] == {}


class TestWhatTheLinkFillsIn:
    def test_the_schedule_comes_from_the_link_not_the_trader(self):
        tag = _tag(_tagging())
        assert sorted(tag["mw_by_hour"]) == list(range(7, 23))
        assert set(tag["mw_by_hour"].values()) == {60.0}  # the smaller side

    def test_it_follows_the_link_when_the_hours_are_edited(self):
        at = _tagging()
        # Re-opening a tag whose link has since been re-allocated must not
        # show the old schedule — the same popup holds both.
        at.session_state["mv_links"][0].mw_by_hour = {7: 5.0}
        at = _emit(_in_dialog(at), {"type": "link_click", "link_id": "L1"})
        at = let_go_of_closed_widgets(at)
        assert _tag(_in_dialog(at).run())["mw_by_hour"] == {7: 5.0}

    def test_the_two_counterparties_name_the_tag_and_its_markets(self):
        tag = _tag(_tagging())
        assert tag["name"] == "AZPS-BPAT"
        assert (tag["source"]["market"], tag["sink"]["market"]) == ("AZPS", "BPAT")

    def test_the_market_path_is_written_in_pse_codes(self):
        # Not AZPS and BPAT, which are the desk's internal market names: a
        # tag's market path is PSE codes, and the mapping supplies them so
        # nobody has to know both.
        assert [row["pse"] for row in _tag(_tagging())["market_path"]] == [
            "APS01", MAG_PSE, "BPAP01"
        ]

    def test_a_link_into_a_market_puts_mag_at_that_end(self):
        at = _tagging(trades=[BUY_LEG], market="SWPW")
        assert [row["pse"] for row in _tag(at)["market_path"]] == ["APS01", MAG_PSE]


class TestWhatTheTraderTypes:
    def test_a_typed_field_reaches_the_tag(self):
        # Into `mv_tags`, which is plain session state — not left in the
        # widget, whose state Streamlit discards on any run that doesn't
        # render it (a closed dialog renders none of these). What comes
        # *back* out of it on reopening can't be checked here: AppTest
        # reads the previous run's widget tree, so a widget that legitimately
        # vanished makes the next run raise. Verified in the browser
        # instead — see PROJECT.md.
        at = _fill(_tagging())
        assert _tag(at)["source"]["gca"] == "GWA"
        assert _tag(at)["sink"]["lca"] == "BPAT"

    def test_a_row_added_to_the_market_path_reaches_the_tag(self):
        at = _fill(_tagging())
        at.session_state["mv_tag_L1_market_path"] = {
            "edited_rows": {}, "deleted_rows": [],
            "added_rows": [{"pse": "EEMU", "product": ""}],
        }
        at = _in_dialog(at).run()
        assert [row["pse"] for row in _tag(at)["market_path"]][-1] == "EEMU"

    def test_the_seed_of_a_table_never_moves_while_it_is_edited(self):
        # The same trap as the schedule grid's double-entry bug: a
        # data_editor's element id is a hash of the data it is handed, so a
        # seed rebuilt from its own output drops the next edit.
        at = _fill(_tagging())
        before = at.session_state["mv_tag_L1_market_path_seed"].copy()
        at.session_state["mv_tag_L1_market_path"] = {
            "edited_rows": {0: {"pse": "TYPED"}}, "added_rows": [], "deleted_rows": [],
        }
        at = _in_dialog(at).run()
        after = at.session_state["mv_tag_L1_market_path_seed"]
        assert after["pse"].tolist() == before["pse"].tolist()
        assert _tag(at)["market_path"][0]["pse"] == "TYPED"


class TestGenerating:
    def test_an_unfinished_tag_cannot_be_generated(self):
        at = _tagging()
        assert any("GCA is required." in e.value for e in at.error)
        assert at.button("mv_tag_L1_generate").disabled

    def test_a_finished_one_can(self, redirect_tags):
        at = _fill(_tagging())
        assert not at.button("mv_tag_L1_generate").disabled
        at = _click(at, "Generate tag file")
        written = list((_day_folder(redirect_tags)).glob("*.xlsx"))
        assert len(written) == 1, [s.value for s in at.error]
        assert any("Tag written to" in s.value for s in at.success)

    def test_what_it_writes_is_the_links_own_schedule(self, redirect_tags):
        _click(_fill(_tagging()), "Generate tag file")
        path = next((_day_folder(redirect_tags)).glob("*.xlsx"))
        ws = openpyxl.load_workbook(path)["BIDS"]
        # PPT HE7-22 at 60 MW, on the sheet's EPT rows (HE10-24 + the first
        # star) — see data/tagfiles/bilateral.py.
        assert [r for r in range(14, 41) if ws.cell(r, 4).value is not None] == list(
            range(23, 39)
        )
        assert ws.cell(23, 4).value == 60.0
        assert ws["D8"].value == "GWA"

    def test_generating_the_same_tag_twice_says_so_rather_than_replacing_it(
        self, redirect_tags
    ):
        at = _click(_fill(_tagging()), "Generate tag file")
        at = _click(_in_dialog(at), "Generate tag file")
        assert any("already exists" in e.value for e in at.error)
        assert len(list((_day_folder(redirect_tags)).glob("*.xlsx"))) == 1

    def test_unless_the_trader_ticks_overwrite(self, redirect_tags):
        at = _click(_fill(_tagging()), "Generate tag file")
        at.checkbox("mv_tag_L1_overwrite").set_value(True)
        at = _click(_in_dialog(at).run(), "Generate tag file")
        assert any("Tag written to" in s.value for s in at.success)
        assert len(list((_day_folder(redirect_tags)).glob("*.xlsx"))) == 1

    def test_a_write_that_fails_is_reported_not_swallowed(self, redirect_tags, monkeypatch):
        def boom(tag, overwrite=False):
            raise RuntimeError("the drive is not mounted")

        monkeypatch.setattr(tagfile, "write_tag_file", boom)
        at = _click(_fill(_tagging()), "Generate tag file")
        assert any("the drive is not mounted" in e.value for e in at.error)


class TestTheReservationFillsInItsPath:
    """Only the number is typed. `PathName` in the desk's OASIS summary is
    character-for-character the string the sheet's Path column wants, so a
    path nobody types is a path nobody can mistype."""

    def test_a_reference_fills_its_path(self):
        at = _type_reservation(_fill(_tagging()), "109267045")
        rows = [r for r in _transmissions(at) if r.get("reservation")]
        assert rows[0]["path"] == "WS/NWMT/NWMT-WAUW/MATL.NWMT-CROSSOVER/"

    def test_the_editor_is_shown_the_filled_path(self):
        # The fill happens in Python, so the widget has to be renamed for
        # the trader to see it — the seed alone would sit under the old diff.
        at = _type_reservation(_fill(_tagging()), "109267045")
        seed = at.session_state["mv_tag_L1_transmissions_seed"]
        assert "WS/NWMT/NWMT-WAUW/MATL.NWMT-CROSSOVER/" in seed["path"].tolist()

    def test_two_rows_each_get_their_own(self):
        at = _type_reservation(_fill(_tagging()), "109267045", "110616287")
        paths = [r["path"] for r in _transmissions(at) if r.get("reservation")]
        assert paths == [
            "WS/NWMT/NWMT-WAUW/MATL.NWMT-CROSSOVER/",
            "W/TEPC/TEPC-AZPS/TEPC.SYS-SAGUARO500/",
        ]

    def test_it_says_what_the_reservation_grants(self):
        # Tagging more MW than the reservation covers is the kind of error
        # that only surfaces days later.
        at = _type_reservation(_fill(_tagging()), "109267045")
        assert any("40 MW granted" in c.value for c in at.caption)

    def test_it_settles_rather_than_filling_on_every_run(self):
        at = _type_reservation(_fill(_tagging()), "109267045")
        before = at.session_state["mv_tag_L1_transmissions_ver"]
        at = _in_dialog(at).run()
        assert at.session_state["mv_tag_L1_transmissions_ver"] == before


class TestWhatItLeavesAlone:
    def test_a_reference_that_is_not_one_is_never_looked_up(self):
        # FCATBTEP goes in the same column and is not an OASIS reference.
        at = _type_reservation(_fill(_tagging()), "FCATBTEP")
        rows = [r for r in _transmissions(at) if r.get("reservation")]
        assert not rows[0].get("path")
        assert not [c for c in at.caption if "FCATBTEP" in c.value]

    def test_a_reservation_that_is_not_ours_leaves_the_path_empty(self):
        at = _type_reservation(_fill(_tagging()), "110752124")
        rows = [r for r in _transmissions(at) if r.get("reservation")]
        assert not rows[0].get("path")

    def test_and_says_so_rather_than_failing_quietly(self):
        at = _type_reservation(_fill(_tagging()), "110752124")
        assert any("isn't one of our own" in c.value for c in at.caption)

    def test_a_path_the_scheduler_typed_is_never_overwritten(self):
        at = _fill(_tagging())
        key = _transmission_key(at)
        at.session_state[key] = {
            "edited_rows": {}, "deleted_rows": [],
            "added_rows": [{"reservation": "109267045", "path": "my own wording"}],
        }
        at = _in_dialog(at).run()
        rows = [r for r in _transmissions(at) if r.get("reservation")]
        assert rows[0]["path"] == "my own wording"

    def test_but_a_path_it_filled_follows_a_changed_reference(self):
        at = _type_reservation(_fill(_tagging()), "109267045")
        key = _transmission_key(at)
        at.session_state[key] = {
            "edited_rows": {0: {"reservation": "110616287"}},
            "added_rows": [], "deleted_rows": [],
        }
        at = _in_dialog(at).run()
        rows = [r for r in _transmissions(at) if r.get("reservation")]
        assert rows[0]["path"] == "W/TEPC/TEPC-AZPS/TEPC.SYS-SAGUARO500/"


class TestThePathString:
    """The path as it stands on chat: generated, displayed to be copied, and
    replaced by whatever comes back. Nothing here reads it — on a link MAG
    doesn't have to tag, carrying the string *is* the job."""

    def test_a_new_link_opens_with_the_only_part_anyone_knows(self):
        at = _drawing()
        assert at.session_state["mv_tags"][PENDING_TAG]["path_string"] == OPENING_PATH

    def test_it_is_shown_where_it_can_be_copied(self):
        # st.code, not st.text: it is the one element that comes with a copy
        # button, and the whole point is pasting this into a chat window.
        assert OPENING_PATH in _shown_path(_tagging())

    def test_pasting_the_next_round_replaces_it(self):
        at = _paste_path(_tagging(), CHAT_PATH)
        assert _tag(at)["path_string"] == CHAT_PATH

    def test_and_the_display_follows_on_that_same_run(self):
        # The entry is read before the display is drawn, so a paste shows up
        # without costing a rerun inside the dialog.
        at = _paste_path(_tagging(), CHAT_PATH)
        assert CHAT_PATH in _shown_path(at)
        assert OPENING_PATH not in _shown_path(at)

    def test_the_box_is_empty_again_for_the_round_after(self):
        at = _paste_path(_tagging(), CHAT_PATH)
        assert at.text_input(_path_entry_key(at)).value == ""

    def test_an_empty_box_leaves_the_path_where_it_was(self):
        at = _paste_path(_tagging(), CHAT_PATH)
        at = _paste_path(at, "   ")
        assert _tag(at)["path_string"] == CHAT_PATH

    def test_it_is_still_there_when_the_link_is_opened_again(self):
        at = _paste_path(_tagging(), CHAT_PATH)
        at = _click(at, "Cancel")
        at = _emit(_in_dialog(at), {"type": "link_click", "link_id": "L1"})
        assert CHAT_PATH in _shown_path(_in_dialog(at).run())

    def test_what_was_pasted_before_the_link_existed_moves_onto_it(self):
        at = _paste_path(_drawing(), CHAT_PATH, link_id=PENDING_TAG)
        at = _click(at, "Create link")
        assert _tag(at)["path_string"] == CHAT_PATH

    def test_a_link_into_a_market_ends_at_mag_rather_than_a_placeholder(self):
        # Selling into SWPW *is* MAG sinking it — there's nobody past that.
        at = _tagging(trades=(BUY_LEG,), frm="session:0", market="SWPW")
        assert "??-APS01-MAG001(s)" in _shown_path(at)


class TestWhereTheTagFieldsSit:
    """Behind an expander, because most links never need them — see
    ui.scheduling.tag's docstring."""

    def test_shut_on_a_link_being_drawn(self):
        assert _tag_expander(_drawing()).proto.expanded is False

    def test_open_on_a_link_being_looked_at_again(self):
        assert _tag_expander(_tagging()).proto.expanded is True

    def test_the_fields_are_reachable_either_way(self):
        # Collapsed is not absent: the widgets render, hold their state and
        # feed the tag exactly as before.
        at = _drawing()
        at.text_input("mv_tag_pending_Source_gca").set_value("GWA")
        at = _in_dialog(at).run()
        assert at.session_state["mv_tags"][PENDING_TAG]["source"]["gca"] == "GWA"

    def test_the_path_string_is_never_behind_them(self):
        at = _tagging()
        assert not [block.value for block in _tag_expander(at).code]


class TestLookingUpOldTags:
    """"Lookup old tags": the button under the path string. Reads OATI
    history (stubbed here — the live wiring is covered by
    tests/data/test_tags_history.py and tests/domain/test_tag_recipes.py)
    and ranks it against what the tag already knows. No path string is read
    anywhere in this — tag_query works off the market path and GCA/LCA."""

    def test_a_shared_counterparty_surfaces_a_route(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_tagging(), "Lookup old tags")
        assert any("GWA" in m.value for m in at.markdown)

    def test_it_works_before_anything_is_typed(self, monkeypatch):
        # The whole point: the counterparties are already known from the
        # market path's own default, before a single field is filled in.
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_drawing(), "Lookup old tags")
        assert not at.exception, [e.value for e in at.exception]
        assert any("GWA" in m.value for m in at.markdown)

    def test_an_unrelated_route_never_surfaces(self, monkeypatch):
        _stub_history(monkeypatch, [UNRELATED_TAG])
        at = _click(_tagging(), "Lookup old tags")
        assert "shares a counterparty" in " ".join(c.value for c in at.caption)

    def test_an_error_is_a_warning_not_a_crash(self, monkeypatch):
        _stub_history(monkeypatch, [], error="no route to host")
        at = _click(_tagging(), "Lookup old tags")
        assert not at.exception, [e.value for e in at.exception]
        assert any("no route to host" in w.value for w in at.warning)

    def test_it_searches_the_trailing_twelve_months(self, monkeypatch):
        captured = {}
        _stub_history(monkeypatch, [], capture=captured)
        _click(_tagging(), "Lookup old tags")
        assert (captured["stop"] - captured["start"]).days == 365


class TestUsingAHistoricalRoute:
    """"Use this route" — display-only until this button is pressed."""

    def _result(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        return _use_route(_click(_tagging(), "Lookup old tags"))

    def test_nothing_is_written_before_the_button_is_pressed(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_tagging(), "Lookup old tags")
        assert _tag(at)["source"]["gca"] == ""

    def test_it_fills_the_market_path(self, monkeypatch):
        at = self._result(monkeypatch)
        assert [row["pse"] for row in _tag(at)["market_path"]] == [
            "APS01", "MAG001", "BPAP01"
        ]

    def test_it_fills_gca_and_lca(self, monkeypatch):
        at = self._result(monkeypatch)
        assert _tag(at)["source"]["gca"] == "GWA"
        assert _tag(at)["sink"]["lca"] == "SWPW"

    def test_it_fills_the_source_and_sink_points(self, monkeypatch):
        at = self._result(monkeypatch)
        assert _tag(at)["source"]["point"] == "RIMROCK"
        assert _tag(at)["sink"]["point"] == "SWPW_HUB"

    def test_it_brings_no_transmission_rows_at_all(self, monkeypatch):
        # A reservation number belongs to a day, not to a route, so last
        # month's are dead on arrival and the rows they came on were being
        # deleted more often than kept.
        at = self._result(monkeypatch)
        assert _numbers(at) == []

    def test_it_leaves_carbon_copy_and_the_name_alone(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _tagging()
        at.text_input("mv_tag_L1_name").set_value("KEEP-ME")
        at = _in_dialog(at).run()
        at = _use_route(_click(at, "Lookup old tags"))
        assert _tag(at)["name"] == "KEEP-ME"
        assert _tag(at)["carbon_copy"] == []

    def test_the_expander_opens_on_a_link_still_being_drawn(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _use_route(_click(_drawing(), "Lookup old tags"))
        assert _tag_expander(at).proto.expanded is True

#: The same route, but sinking into CISO rather than SWPW — what an
#: EPE-MAG-SWPW link kept being offered before a market end settled its LCA.
INTO_CISO = {
    "tag_index": 3,
    "gca": "GWA",
    "lca": "CISO",
    "flow_date": date(2026, 9, 1),
    "market_path": [
        {"pse": "APS01", "product": "G-F", "contracts": ""},
        {"pse": "MAG001", "product": "L", "contracts": ""},
    ],
    "physical_path": [
        {"type": "G", "tp": "", "por_name": "PALOVERDE", "por_ca": "GWA",
         "pod_name": "", "pod_ca": "", "reservations": []},
        {"type": "L", "tp": "", "por_name": "", "por_ca": "",
         "pod_name": "CISO_HUB", "pod_ca": "CISO", "reservations": []},
    ],
}


class TestWhatARouteLeavesForTheScheduler:
    """A route is a shape, not a day's paperwork."""

    def _result(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        return _use_route(_click(_tagging(), "Lookup old tags"))

    def test_the_transmission_block_is_left_empty(self, monkeypatch):
        assert _numbers(self._result(monkeypatch)) == []

    def test_even_the_reference_that_is_still_good_today(self, monkeypatch):
        # 109267045 covers this flow date and is still not carried: the
        # rule is that the number is typed, full stop.
        assert "109267045" not in _numbers(self._result(monkeypatch))

    def test_and_no_path_arrives_without_one(self, monkeypatch):
        # The invariant: every path on screen has a reference behind it
        # that a scheduler can check.
        for row in _transmissions(self._result(monkeypatch)):
            assert not (row["path"] and not row["reservation"])

    def test_a_number_typed_afterwards_still_fills_its_own_path(self, monkeypatch):
        # Which is the one direction that puts nothing on screen nobody
        # asked for.
        at = _type_reservation(self._result(monkeypatch), "109267045")
        assert _transmissions(at)[0]["path"] == OASIS["109267045"]["path"]


class TestALinkIntoAMarket:
    """A market end is where MAG sinks or sources the power, so its control
    area is settled — not a field the lookup may guess around."""

    def _into_swpw(self, monkeypatch, history):
        _stub_history(monkeypatch, history)
        at = _tagging(trades=(BUY_LEG,), frm="session:0", market="SWPW")
        return _click(at, "Lookup old tags")

    def test_a_route_sinking_somewhere_else_is_not_offered(self, monkeypatch):
        at = self._into_swpw(monkeypatch, [INTO_CISO])
        assert "shares a counterparty" in " ".join(c.value for c in at.caption)

    def test_while_one_sinking_into_that_market_is(self, monkeypatch):
        into_swpw = dict(INTO_CISO, tag_index=4, lca="SWPW")
        into_swpw["physical_path"] = [
            dict(INTO_CISO["physical_path"][0]),
            {"type": "L", "tp": "", "por_name": "", "por_ca": "",
             "pod_name": "SWPW_HUB", "pod_ca": "SWPW", "reservations": []},
        ]
        at = self._into_swpw(monkeypatch, [INTO_CISO, into_swpw])
        shown = " ".join(m.value for m in at.markdown)
        assert "SWPW_HUB" in shown
        assert "CISO_HUB" not in shown


class TestAnEditorNeverOutlivesItsOwnHeight:
    """Why a deleted row used to come back the moment anything else was
    touched.

    An editor is sized to the rows it holds, and `height` is one of the
    properties Streamlit hashes into a data editor's element id. So the run
    *after* a deletion handed the widget a new id, looked the trader's diff
    up under a name that held nothing, and returned the old seed untouched.
    Traced in Chrome against the live app:

        run 1  in=3  state={deleted_rows: [1]}  out=2
        run 2  in=2  state={deleted_rows: [1]}  out=3

    Tested here as the rule itself rather than through the page: AppTest
    carries no data editor state between runs at all, so the editor is
    remounted on every run there and neither the bug nor the fix is visible
    — a page-level test would pass either way. The browser is where this
    one is checked; see PROJECT.md.
    """

    def test_the_seed_moves_whenever_the_height_does(self):
        for count in range(0, 8):
            before = [{"reservation": str(i)} for i in range(count)]
            after = before + [{"reservation": "new"}]
            if table_height(before) == table_height(after):
                continue  # under MIN_TABLE_ROWS the editor doesn't resize
            assert seed_is_stale(pd.DataFrame(before), after), count

    def test_and_a_row_deleted_moves_it_the_same_way(self):
        seed = pd.DataFrame([{"reservation": "a"}, {"reservation": "b"}])
        assert seed_is_stale(seed, [{"reservation": "a"}])

    def test_but_typing_within_a_row_leaves_it_alone(self):
        # The other half of the rule: a seed that moved on every keystroke
        # would rename the widget mid-edit and drop the *next* edit.
        seed = pd.DataFrame([{"reservation": "a"}, {"reservation": "b"}])
        assert not seed_is_stale(
            seed, [{"reservation": "109267045"}, {"reservation": "b"}]
        )


class TestADeletionThatHasToStick:
    def test_a_path_cell_emptied_by_hand_stays_empty(self):
        # The lookup fills a blank Path once. After that the cell is the
        # scheduler's: a path that comes back on the next rerun is not a
        # lookup being helpful, it is a lookup overruling someone.
        at = _type_reservation(_fill(_tagging()), "109267045")
        assert _transmissions(at)[0]["path"]
        at.session_state[_transmission_key(at)] = {
            "edited_rows": {0: {"path": ""}}, "deleted_rows": [], "added_rows": [],
        }
        at = _in_dialog(at).run()
        assert _transmissions(at)[0]["path"] == ""
        # And stays empty through whatever is touched next.
        at = _fill(at, gca="GWA2")
        assert _transmissions(at)[0]["path"] == ""


class TestSplittingALinkAcrossSeveralTags:
    """One schedule, several paths, several tags. The default is still one —
    a link opens the way it always did — and the thing that has to hold
    however it is split is that the tags add up to the link."""

    def _edit(self, mw, hours=range(7, 23)):
        """One tag's share, the way Streamlit stores an edit to the hour
        grid — see test_multiday_links.py, same idiom."""
        return {
            "edited_rows": {0: {str(h): float(mw) for h in hours}},
            "added_rows": [], "deleted_rows": [],
        }

    def _hours(self, at, tag_key, mw, hours=range(7, 23)):
        at.session_state[f"mv_tag_{tag_key}_hours"] = self._edit(mw, hours)
        return _in_dialog(at).run()

    def test_a_link_opens_with_one_tag_and_no_tabs(self):
        at = _tagging()
        assert list(at.session_state["mv_tags"]) == ["L1"]
        assert not at.tabs

    def test_splitting_gives_it_a_second(self):
        at = _click(_tagging(), "Split")
        assert list(at.session_state["mv_tags"]) == ["L1", "L1_t2"]
        assert [tab.label for tab in at.tabs] == ["Tag 1", "Tag 2"]

    def test_the_new_one_starts_with_no_hours_of_its_own(self):
        # So the split still adds up the moment it is made: everything is
        # on the tag that was already there, and the trader moves hours
        # across rather than finding the link over-tagged on the way in.
        at = _click(_tagging(), "Split")
        assert _tag(at, "L1_t2")["mw_by_hour"] == {}
        assert _tag(at)["mw_by_hour"]
        # And the split adds up on the way in, so the check says nothing.
        assert not [e for e in at.error if "carry" in e.value]

    def test_and_a_name_that_does_not_collide_with_the_first(self):
        # Two tags of one link would otherwise be written to one file.
        at = _click(_tagging(), "Split")
        assert _tag(at)["name"] != _tag(at, "L1_t2")["name"]

    def test_each_tag_carries_its_own_path_string(self):
        at = _click(_tagging(), "Split")
        at = _paste_path(at, "??-APS01-MAG001-SOMEONE-??", link_id="L1_t2")
        assert _tag(at, "L1_t2")["path_string"] == "??-APS01-MAG001-SOMEONE-??"
        assert _tag(at)["path_string"] == OPENING_PATH

    def test_and_its_own_fields(self):
        at = _click(_tagging(), "Split")
        at.text_input("mv_tag_L1_t2_Source_gca").set_value("GWA")
        at = _in_dialog(at).run()
        assert _tag(at, "L1_t2")["source"]["gca"] == "GWA"
        assert _tag(at)["source"]["gca"] == ""

    def test_moving_hours_onto_the_second_tag_leaves_the_link_over_tagged(self):
        # Until the first one gives them up — which is the whole point of
        # the check.
        at = self._hours(_click(_tagging(), "Split"), "L1_t2", 60)
        assert any("more than this link" in e.value for e in at.error)

    def test_and_taking_too_many_off_the_first_leaves_it_short(self):
        at = _click(_tagging(), "Split")
        at = self._hours(at, "L1", 10)
        assert any("less than this link" in e.value for e in at.error)

    def test_a_split_that_adds_up_says_nothing(self):
        # Both shares in one run: AppTest carries no data editor state
        # between runs, so typing into one grid and then the other would
        # lose the first edit here — a limitation of the harness, not of
        # the page (see TestAnEditorNeverOutlivesItsOwnHeight).
        at = _click(_tagging(), "Split")
        at.session_state["mv_tag_L1_hours"] = self._edit(20)
        at.session_state["mv_tag_L1_t2_hours"] = self._edit(40)
        at = _in_dialog(at).run()          # the link carries 60
        assert not [e for e in at.error if "carry" in e.value]

    def _fill(self, at, tag_key, mw):
        """Type one MW into that tab's fill box and press its own Fill —
        every tab has one, so the button has to be picked by key."""
        at.number_input(f"mv_tag_{tag_key}_fill").set_value(mw)
        at = _in_dialog(at).run()
        key = f"mv_tag_{tag_key}_fill_go"
        [b for b in _in_dialog(at).button if b.key == key][0].click().run()
        return _in_dialog(at).run()

    def test_one_mw_can_be_filled_across_every_hour_at_once(self):
        # Most tags carry the same MW all day, and typing it into twenty-four
        # cells is twenty-four chances to fat-finger one.
        at = self._fill(_click(_tagging(), "Split"), "L1_t2", 15.0)
        assert _tag(at, "L1_t2")["mw_by_hour"] == {h: 15.0 for h in range(7, 23)}

    def test_and_only_across_the_hours_the_link_runs(self):
        # An hour the link doesn't run is an hour this tag can't have, and
        # filling it would only hand the trader a check to clear.
        at = self._fill(_click(_tagging(), "Split"), "L1_t2", 15.0)
        assert set(_tag(at, "L1_t2")["mw_by_hour"]) == set(_tag(at)["mw_by_hour"])

    def test_the_grid_shows_what_was_filled(self):
        # A fill the trader can't see is a fill they'll do again by hand.
        at = self._fill(_click(_tagging(), "Split"), "L1_t2", 15.0)
        assert at.session_state["mv_tag_L1_t2_hours_seed"] == {
            h: 15.0 for h in range(7, 23)
        }
        assert at.session_state["mv_tag_L1_t2_hours_ver"] == 1

    def test_nothing_is_filled_until_a_number_is_typed(self):
        at = _click(_tagging(), "Split")
        assert [b for b in at.button if b.label == "Fill"][0].disabled

    def test_removing_a_tag_puts_the_link_back_the_way_it_was(self):
        at = _click(_tagging(), "Split")
        at = _click(at, "Remove")
        assert list(at.session_state["mv_tags"]) == ["L1"]
        assert not at.tabs

    def test_the_first_tag_can_never_be_removed(self):
        # It is keyed by the link itself, and a link always has a tag.
        at = _click(_tagging(), "Split")
        assert len([b for b in at.button if b.label == "Remove"]) == 1

    def test_a_link_still_being_drawn_can_be_split_too(self):
        at = _click(_drawing(), "Split")
        assert sorted(at.session_state["mv_tags"]) == ["pending", "pending_t2"]

    def test_and_both_move_onto_the_link_when_it_is_created(self):
        at = _click(_drawing(), "Split")
        at = _click(at, "Create link")
        assert sorted(at.session_state["mv_tags"]) == ["L1", "L1_t2"]

    def test_abandoning_a_split_link_takes_every_tag_with_it(self):
        at = _click(_drawing(), "Split")
        at = _click(at, "Cancel")
        assert at.session_state["mv_tags"] == {}

    def test_deleting_the_link_does_too(self):
        at = _click(_tagging(), "Split")
        at = _click(at, "Delete")
        assert at.session_state["mv_tags"] == {}

    def test_one_state_covers_every_tag_of_a_link(self):
        # The board draws one line per link, and a line is one colour.
        at = _click(_tagging(), "Split")
        at.radio("mv_tag_L1_state").set_value("tagged")
        at = _in_dialog(at).run()
        assert _tag(at)["state"] == "tagged"
        assert _tag(at, "L1_t2")["state"] == "tagged"


class TestGeneratingASplitLinksTags:
    """One file per press, holding the tags marked ready. Marking one and
    generating is how a single tag of a split goes out on its own."""

    def _edit(self, mw):
        return {
            "edited_rows": {0: {str(h): float(mw) for h in range(7, 23)}},
            "added_rows": [], "deleted_rows": [],
        }

    def _ready(self, at, *keys, hours=None):
        """Fill in what a tag needs and tick it ready — every field in one
        run, since AppTest carries no data editor state between them."""
        for key in keys:
            at.text_input(f"mv_tag_{key}_Source_gca").set_value("GWA")
            at.text_input(f"mv_tag_{key}_Sink_lca").set_value("BPAT")
            at.checkbox(f"mv_tag_{key}_ready").set_value(True)
            if hours:
                at.session_state[f"mv_tag_{key}_hours"] = self._edit(hours[key])
        return _in_dialog(at).run()

    def _files(self, folder):
        return sorted(p.name for p in folder.rglob("*.xlsx"))

    def test_a_new_split_tag_is_not_ready_until_it_is_said_to_be(self):
        # Or it would hold up the file for the ones that are finished.
        at = _click(_tagging(), "Split")
        assert _tag(at, "L1_t2")["ready"] is False
        assert _tag(at)["ready"] is True

    def test_one_ready_tag_writes_one_ordinary_file(self, redirect_tags):
        at = _click(_tagging(), "Split")           # tag 2 is not ready
        at = self._ready(at, "L1")
        at = _click(at, "Generate tag file")
        assert len(self._files(redirect_tags)) == 1
        assert "SPLIT" not in self._files(redirect_tags)[0]

    def test_an_unready_tag_does_not_hold_up_the_ready_one(self, redirect_tags):
        # Tag 2 has no hours, no GCA and no LCA, and none of that counts.
        at = self._ready(_click(_tagging(), "Split"), "L1")
        assert not [b for b in at.button if b.label.startswith("Generate")][0].disabled

    def test_two_ready_tags_share_one_file(self, redirect_tags):
        at = _click(_tagging(), "Split")
        at = self._ready(at, "L1", "L1_t2", hours={"L1": 20, "L1_t2": 40})
        at = _click(at, "Generate tag file")
        written = self._files(redirect_tags)
        assert len(written) == 1
        assert written[0].endswith("SPLIT.xlsx")

    def test_and_both_are_in_it(self, redirect_tags):
        at = _click(_tagging(), "Split")
        at = self._ready(at, "L1", "L1_t2", hours={"L1": 20, "L1_t2": 40})
        _click(at, "Generate tag file")
        import openpyxl

        ws = openpyxl.load_workbook(
            next(redirect_tags.rglob("*.xlsx"))
        )["BIDS"]
        assert ws["D8"].value == "GWA" and ws["G8"].value == "GWA"

    def test_nothing_is_written_when_nothing_is_ready(self, redirect_tags):
        at = _click(_tagging(), "Split")
        at.checkbox("mv_tag_L1_ready").set_value(False)
        at = _in_dialog(at).run()
        assert [b for b in at.button if b.label.startswith("Generate")][0].disabled

    def test_a_split_that_does_not_add_up_is_not_written_at_all(self):
        # Both tags complete in themselves and both ready; the only thing
        # wrong is that 20 and 30 are not the 60 on the link.
        at = _click(_tagging(), "Split")
        at = self._ready(at, "L1", "L1_t2", hours={"L1": 20, "L1_t2": 30})
        assert any("less than this link" in e.value for e in at.error)
        assert [b for b in at.button if b.label.startswith("Generate")][0].disabled

    def test_but_it_does_not_stop_a_partial_split_going_out(self):
        # With tag 2 still unready the shares obviously don't add up yet,
        # and that is exactly when a trader writes the part that is done.
        at = self._ready(_click(_tagging(), "Split"), "L1", hours={"L1": 20})
        assert any("less than this link" in e.value for e in at.error)
        assert not [b for b in at.button if b.label.startswith("Generate")][0].disabled

    def test_an_unsplit_link_still_writes_the_one(self, redirect_tags):
        at = _tagging()
        at.text_input("mv_tag_L1_Source_gca").set_value("GWA")
        at.text_input("mv_tag_L1_Sink_lca").set_value("BPAT")
        at = _in_dialog(at).run()
        at = _click(at, "Generate tag file")
        assert len(self._files(redirect_tags)) == 1


class TestWhoStillHasToTagThisLink:
    """Three states, set by hand in the popup and read by the board, which
    colours each link's line by it. Nothing infers them yet."""

    def _radio(self, at, link_id="L1"):
        return at.radio(f"mv_tag_{link_id}_state")

    def test_a_link_opens_as_one_mag_still_has_to_tag(self):
        at = _tagging()
        assert _tag(at)["state"] == "mag"
        assert self._radio(at).value == "mag"

    def test_the_three_are_offered_in_words(self):
        at = _tagging()
        assert self._radio(at).options == [
            "MAG tags it", "Counterparty tags it", "Tagged"
        ]

    def test_choosing_one_lands_on_the_tag(self):
        at = _tagging()
        self._radio(at).set_value("tagged")
        at = _in_dialog(at).run()
        assert _tag(at)["state"] == "tagged"

    def test_a_link_still_being_drawn_carries_its_state_onto_itself(self):
        at = _drawing()
        at.radio(f"mv_tag_{PENDING_TAG}_state").set_value("counterparty")
        at = _in_dialog(at).run()
        at = _click(at, "Create link")
        assert _tag(at)["state"] == "counterparty"


class TestWhatSinkingIntoSwpwFillsIn:
    """The desk's own convention for a link sinking into SWPW: SPP in the
    Market cell, SWPW as the LCA, the delivery point named in the header
    and the Sink cell left empty. See domain.tags.MARKET_SINK_FIELDS."""

    def _into_swpw(self):
        return _tagging(trades=(BUY_LEG,), frm="session:0", market="SWPW")

    def test_the_header_names_the_delivery_point(self):
        assert _tag(self._into_swpw())["label"] == "SOURCE-SWPW(??)"

    def test_the_market_cell_is_the_pool_and_the_lca_the_market(self):
        tag = _tag(self._into_swpw())
        assert tag["sink"]["market"] == "SPP"
        assert tag["sink"]["lca"] == "SWPW"

    def test_and_the_sink_cell_is_left_for_nobody(self):
        at = self._into_swpw()
        assert _tag(at)["sink"]["point"] == ""
        assert at.text_input("mv_tag_L1_Sink_point").value == ""

    def test_an_old_route_does_not_fill_it_either(self, monkeypatch):
        # The route is a guess; the market end is not.
        into_swpw = dict(HISTORICAL_TAG, lca="SWPW")
        into_swpw["market_path"] = [
            {"pse": "APS01", "product": "G-F", "contracts": ""},
            {"pse": "MAG001", "product": "L", "contracts": ""},
        ]
        _stub_history(monkeypatch, [into_swpw])
        at = _use_route(_click(self._into_swpw(), "Lookup old tags"))
        assert _tag(at)["sink"]["point"] == ""
        # The rest of the route still landed.
        assert _tag(at)["source"]["point"] == "RIMROCK"


def _many_routes(count):
    """`count` distinct recipes, all sharing this link's counterparties so
    every one of them is offered. They differ only in where the power
    starts, which is enough to make each its own route."""
    routes = []
    for n in range(count):
        route = copy.deepcopy(HISTORICAL_TAG)
        route["tag_index"] = 100 + n
        route["physical_path"][0]["por_name"] = f"PLANT{n}"
        routes.append(route)
    return routes


class TestSeeingMoreThanTheFirstFour:
    """The ranking is a guess, and a guess that puts the right route fifth
    is no use if fifth can't be reached."""

    def _routes(self, at):
        return [b for b in at.button if b.label == "Use this route"]

    def _turn(self, at, label):
        """Press a pager button and let the tree settle. The press reruns
        the dialog fragment, and AppTest needs one more run after that
        before its element tree catches up — the session state is right
        immediately, the widgets are a render behind."""
        return _in_dialog(_click(at, label)).run()

    def _lookup(self, monkeypatch, count=9):
        _stub_history(monkeypatch, _many_routes(count))
        return _click(_tagging(), "Lookup old tags")

    def test_four_at_a_time(self, monkeypatch):
        at = self._lookup(monkeypatch)
        assert len(self._routes(at)) == 4

    def test_and_it_says_which_four_of_how_many(self, monkeypatch):
        at = self._lookup(monkeypatch)
        assert any("1-4 of 9" in c.value for c in at.caption)

    def test_next_shows_the_four_behind_them(self, monkeypatch):
        at = self._lookup(monkeypatch)
        first = [b.key for b in self._routes(at)]
        at = self._turn(at, "Next 4")
        assert [b.key for b in self._routes(at)] != first
        assert any("5-8 of 9" in c.value for c in at.caption)

    def test_and_back_returns(self, monkeypatch):
        at = self._lookup(monkeypatch)
        first = [b.key for b in self._routes(at)]
        at = self._turn(self._turn(at, "Next 4"), "Back")
        assert [b.key for b in self._routes(at)] == first

    def test_the_last_page_can_be_short(self, monkeypatch):
        at = self._lookup(monkeypatch)
        at = self._turn(self._turn(at, "Next 4"), "Next 4")
        assert len(self._routes(at)) == 1
        assert any("9-9 of 9" in c.value for c in at.caption)

    def test_there_is_nowhere_to_go_from_the_ends(self, monkeypatch):
        at = self._lookup(monkeypatch)
        assert [b for b in at.button if b.label == "Back"][0].disabled
        at = self._turn(self._turn(at, "Next 4"), "Next 4")
        assert [b for b in at.button if b.label == "Next 4"][0].disabled

    def test_no_pager_at_all_when_four_is_the_lot(self, monkeypatch):
        at = self._lookup(monkeypatch, count=4)
        assert not [b for b in at.button if b.label == "Next 4"]

    def test_a_fresh_lookup_starts_from_the_top(self, monkeypatch):
        at = self._lookup(monkeypatch)
        at = self._turn(at, "Next 4")
        at = _click(at, "Lookup old tags")
        assert any("1-4 of 9" in c.value for c in at.caption)


class TestWhatARouteCardShows:
    def test_where_the_power_starts_and_ends(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_tagging(), "Lookup old tags")
        assert any(
            "RIMROCK (GWA) → SWPW_HUB (SWPW)" in m.value for m in at.markdown
        )

    def test_the_wires_it_crosses(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_tagging(), "Lookup old tags")
        said = " ".join(c.value for c in at.caption)
        assert "NWMT MATL.NWMT>CROSSOVER" in said
        assert "TEPC TEPC.SYS>SAGUARO500" in said

    def test_and_how_often_the_desk_has_run_it(self, monkeypatch):
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_tagging(), "Lookup old tags")
        assert any("GWA>SWPW" in m.value for m in at.markdown)
        assert any("used 1×" in c.value for c in at.caption)

    def test_the_pse_chain_it_was_run_under(self, monkeypatch):
        # Without this two routes differing only in a firm/non-firm energy
        # product are two cards nobody can tell apart.
        _stub_history(monkeypatch, [HISTORICAL_TAG])
        at = _click(_tagging(), "Lookup old tags")
        assert any(
            "APS01 (G-F) > MAG001 > BPAP01 (L)" in c.value for c in at.caption
        )


class TestThePopupStaysOnOneScreen:
    """The popup is packed to fit without scrolling — see
    ui.scheduling.links._TIGHTEN. Height itself is a browser question and
    is checked there; what's checked here is the rule that made the last
    three attempts at it look tighter and render broken."""

    def _tighten(self):
        from ui.scheduling.links import _TIGHTEN

        return _TIGHTEN

    @pytest.mark.parametrize("zeroed", ["margin:0", "margin-bottom:0", "min-height:0"])
    def test_no_text_element_has_its_spacing_zeroed(self, zeroed):
        # Streamlit sizes a label, caption or markdown container from
        # exactly these, so zeroing one collapses the container to nothing
        # while its text still paints — it lands on top of the next
        # element. It measures *shorter*, which is the trap.
        for line in self._tighten().splitlines():
            if zeroed not in line.replace(" ", ""):
                continue
            assert not any(
                text in line
                for text in ("WidgetLabel", "stCaptionContainer", "stMarkdownContainer")
            ), line

    def test_generate_sits_in_the_action_row_not_the_tag_fields(self):
        # It costs the tag column nothing down there, and the action row
        # had the space — see ui.scheduling.tag.render_generate.
        at = _tagging()
        assert [b for b in at.button if b.label == "Generate tag file"]
        assert not [
            b for b in _tag_expander(at).button if b.label == "Generate tag file"
        ]

    def test_but_what_stops_it_being_written_stays_with_the_fields(self):
        at = _tagging()
        assert [e for e in _tag_expander(at).error if "GCA is required" in e.value]


class TestAnEditorShowsEveryRowItHolds:
    """A wheel you can't see is a wheel you forget to check, so the
    editors grow with their rows and the popup grows with them — rather
    than hiding the last one inside a fixed frame."""

    def _height(self, rows):
        from ui.scheduling.tag import table_height

        return table_height([{"a": 1}] * rows)

    def test_a_long_block_is_taller_than_a_short_one(self):
        assert self._height(8) > self._height(4) > self._height(3)

    def test_each_extra_row_costs_exactly_one_row(self):
        from ui.scheduling.tag import ROW_HEIGHT

        assert self._height(6) - self._height(5) == ROW_HEIGHT

    def test_it_leaves_room_for_the_blank_row_you_add_with(self):
        from ui.scheduling.tag import HEADER_HEIGHT, ROW_HEIGHT

        # Header, four rows, and the one a dynamic editor keeps at the end.
        assert self._height(4) == HEADER_HEIGHT + 5 * ROW_HEIGHT

    def test_an_empty_block_is_still_worth_typing_into(self):
        from ui.scheduling.tag import MIN_TABLE_ROWS

        assert self._height(0) == self._height(MIN_TABLE_ROWS - 1)

    def test_the_transmission_block_grows_with_what_is_typed_into_it(self):
        from ui.scheduling.tag import table_height

        at = _tagging()
        before = table_height(_transmissions(at))
        at = _type_reservation(at, "109267045", "110616287", "FCATBTEP")
        assert len(_transmissions(at)) == 3
        assert table_height(_transmissions(at)) > before



class TestTheOasisLookupRunsBeforeTheEditorDraws:
    """A guard for a trap no AppTest can see.

    Filling a path after the editor has drawn needs a rerun to show, and a
    rerun on the run that *opens the tag expander* strands every field in
    it: those inputs mount for the first time, the run is thrown away, and
    the browser keeps the empty boxes while the server believes they hold
    the route's values. Server and screen then disagree permanently —
    which AppTest cannot reproduce, because AppTest *is* the server and
    its state was right the whole time. Checked in Chrome, kept here.
    """

    def _render_tag_source(self):
        import inspect

        from ui.scheduling.tag import render_tag

        return inspect.getsource(render_tag)

    def test_a_path_is_looked_up_before_the_transmission_editor(self):
        source = self._render_tag_source()
        first_lookup = source.index("_fill_paths_from_oasis(")
        editor = source.index('"transmissions",')
        assert first_lookup < editor, (
            "the OASIS lookup must run before the transmission editor draws, "
            "or applying a route strands the fields beside it"
        )

    def test_and_again_after_it_for_a_number_just_typed(self):
        # The second call is what catches a reservation the scheduler
        # typed, which only reaches the tag as the editor returns.
        assert self._render_tag_source().count("_fill_paths_from_oasis(") == 2
