"""Writing a bilateral e-Tag workbook.

Unlike the SWPW bid file — whose column layout the desk has evolved by hand,
and which is therefore built from scratch — every one of the 2,638 real tag
files under `TAG_FOLDER` puts the same label in the same cell. The row
numbers below were read off that whole set, not off one example: the
Market/GCA/Source block, the MW columns, the Path block and the Market Path
Product block are identical in all of them, and the Carbon Copy block is
identical wherever it appears at all (it is absent from the older ones).

So this fills a **blank copy of the desk's own template** rather than
recreating it. `bilateral_tag_template.xlsx` beside this module is
`Templates/Tag_dynamique_bilateral - Template - SWPW-SWPP.xlsx` with every
input cell emptied and nothing else touched — its labels, fills, borders,
merges, column widths, zoom and the two cell comments on the `1*`/`2*` rows
all come along. Regenerate it the same way if the desk's template changes.

**The sheet's hour column is EPT; the app is PPT**, exactly as in the SWPW
bid file, and for the same reason the same three `1*`/`2*`/`3*` rows sit
under HE24 in both — see `domain.bidfiles.ppt_to_ept`. Confirmed against the
real files twice over: "… TEPC HE7-14" fills rows 23-30, which is EPT
HE10-17, and a full-day position fills EPT HE4-24 plus the three stars.
"""

from pathlib import Path

from openpyxl import load_workbook

from domain.bidfiles import ppt_to_ept
from domain.tags import (
    carbon_copy_rows,
    day_folder_name,
    folder_covers,
    market_path_rows,
    tag_filename,
    transmission_rows,
)

TAG_FOLDER = Path(r"Y:\West\TAG Bilateral WEST")

#: TEMPORARY, while this feature is under review — every generated tag goes
#: into the desk's own "Test" folder instead of the flow date's real one, so
#: nothing lands where a scheduler would pick it up as a real tag. Flip to
#: False once ready to go live; nothing else about the file changes.
TEST_MODE = True
TEST_SUBFOLDER = "Test"

TEMPLATE = Path(__file__).with_name("bilateral_tag_template.xlsx")
SHEET = "BIDS"

DATE_CELL, LABEL_CELL = "C5", "D5"

#: The six cells of each half, in the order the sheet lists them. The
#: source half's second row is the GCA and the sink half's is the LCA —
#: the only difference between the two blocks.
SOURCE_CELLS = {
    "market": "D7", "gca": "D8", "point": "D9", "pse": "D10",
    "comment": "D11", "contract": "D12",
}
SINK_CELLS = {
    "market": "D44", "lca": "D45", "point": "D46", "pse": "D47",
    "comment": "D48", "contract": "D49",
}

#: Each MW block: the row EPT HE1 sits on, then the three rows holding HE1-3
#: of the *next* EPT day (the last three hours of the PPT day).
SOURCE_HOURS_ROW, SOURCE_STAR_ROWS = 14, (38, 39, 40)
SINK_HOURS_ROW, SINK_STAR_ROWS = 51, (75, 76, 77)

MW_COL, PRICE_COL = 4, 5  # D, E

#: The three list blocks, each running from its first row up to the header
#: of the next block. Their lengths are what domain.tags checks against.
TRANSMISSION_ROW, TRANSMISSION_FIELDS = 80, ("path", "reservation", "mw")
MARKET_PATH_ROW, MARKET_PATH_FIELDS = 94, ("pse", "product", "contract")
CARBON_COPY_ROW, CARBON_COPY_FIELDS = 105, ("type", "market")


def hour_row(base_row, star_rows, he_ppt):
    """The row a PPT hour-ending lands on in a block starting at
    `base_row` — see the module docstring on EPT."""
    he_ept, next_day = ppt_to_ept(he_ppt)
    if next_day:
        return star_rows[he_ept - 1]
    return base_row + he_ept - 1


def _as_number(text):
    """An OASIS reservation number written back as the integer it is in
    every real file, or left as typed when it isn't one (a contract
    reference like 'FCATBTEP' goes in the same column)."""
    try:
        return int(str(text).strip())
    except (TypeError, ValueError):
        return text


def _write_block(ws, first_row, rows, fields, converters=None):
    converters = converters or {}
    for offset, row in enumerate(rows):
        for column, field in enumerate(fields, start=MW_COL):
            value = row.get(field, "")
            if not value:
                continue
            convert = converters.get(field)
            ws.cell(first_row + offset, column).value = (
                convert(value) if convert else value
            )


def _write_hours(ws, base_row, star_rows, mw_by_hour):
    for he_ppt, mw in sorted(mw_by_hour.items()):
        if not mw:
            continue
        ws.cell(hour_row(base_row, star_rows, int(he_ppt)), MW_COL).value = float(mw)


def build_workbook(tag):
    """The template filled in for one tag. The caller validates first —
    `domain.tags.tag_errors` — this only writes what it's given."""
    wb = load_workbook(TEMPLATE)
    ws = wb[SHEET]

    ws[DATE_CELL] = tag["flow_date"]
    ws[LABEL_CELL] = tag.get("label", "")

    for half, cells in ((tag["source"], SOURCE_CELLS), (tag["sink"], SINK_CELLS)):
        for field, cell in cells.items():
            if half.get(field):
                ws[cell] = half[field]

    # Both halves of the sheet carry the same schedule: it is one flow of
    # energy described from each end.
    _write_hours(ws, SOURCE_HOURS_ROW, SOURCE_STAR_ROWS, tag["mw_by_hour"])
    _write_hours(ws, SINK_HOURS_ROW, SINK_STAR_ROWS, tag["mw_by_hour"])

    _write_block(
        ws, TRANSMISSION_ROW, transmission_rows(tag), TRANSMISSION_FIELDS,
        {"reservation": _as_number, "mw": _as_number},
    )
    _write_block(ws, MARKET_PATH_ROW, market_path_rows(tag), MARKET_PATH_FIELDS)
    _write_block(ws, CARBON_COPY_ROW, carbon_copy_rows(tag), CARBON_COPY_FIELDS)
    return wb


def day_folder(flow_date):
    """The folder this flow date's tags belong in.

    The desk files a day under "25 September 2026", or under a two-day
    "25-26 September 2026" when a weekend or a holiday is traded together —
    which is a judgement about the trading session, so an existing folder
    is always preferred and a two-day one is never invented. An unreadable
    root falls through to the single-day name and lets the write itself
    report the problem, which is where the useful message is.
    """
    if TEST_MODE:
        return TAG_FOLDER / TEST_SUBFOLDER
    try:
        for child in sorted(TAG_FOLDER.iterdir()):
            if child.is_dir() and folder_covers(child.name, flow_date):
                return child
    except OSError:
        pass
    return TAG_FOLDER / day_folder_name(flow_date)


def target_path(tag):
    return day_folder(tag["flow_date"]) / tag_filename(
        tag["flow_date"], tag.get("name", "")
    )


def write_tag_file(tag, overwrite=False):
    """Build and save the tag. Returns the path written to.

    Raises FileExistsError if the target is already there and `overwrite`
    isn't set: two tags on one day that happen to share a name are far more
    likely to be a second leg needing its own name than a file anyone
    wanted replaced.
    """
    path = target_path(tag)
    if path.exists() and not overwrite:
        raise FileExistsError(f"{path} already exists.")
    wb = build_workbook(tag)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except OSError as e:
        raise RuntimeError(
            f"Could not write {path.name}: {e}. If that file is open in "
            "Excel, close it and generate again."
        ) from e
    return path
