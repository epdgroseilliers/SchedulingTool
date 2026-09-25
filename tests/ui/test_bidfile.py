"""ui/scheduling/bidfile.py — the SWPW bid-file builder popup, through the
Scheduling View page.

The board is a custom component, so a chip click reaches Python exactly the
way any board event does: `session_state["mv_board"] = {...}` — see
tests/ui/test_scheduling_view.py's own docstring. Every write goes to
tmp_path via the `redirect_swpw` fixture — never Z:\\.
"""

import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

import data.bidfiles.swpw as swpw
from domain.bidfiles import LONG, SHORT
from ui.scheduling.state import BIDFILE_DIALOG, LINK_DIALOG

PAGE_PATH = str(Path(__file__).resolve().parents[2] / "pages" / "1_Scheduling_View.py")
FLOW = date.today() + timedelta(days=1)


@pytest.fixture
def no_bilateral_db(monkeypatch):
    monkeypatch.setattr(
        "ui.scheduling.state.load_trades_for_flow_date", lambda flow_date: ([], None)
    )


@pytest.fixture
def redirect_swpw(tmp_path, monkeypatch):
    monkeypatch.setattr(swpw, "BID_FOLDER", tmp_path)
    monkeypatch.setattr(swpw, "TEST_MODE", True)
    monkeypatch.setattr(swpw, "TEST_FILENAME", "test-output.xlsm")
    # write_bid_file needs a VBA-shell donor to produce a genuine
    # macro-enabled file (see data/bidfiles/swpw.py's _save_as_macro_enabled)
    # — a plain workbook with a dummy xl/vbaProject.bin entry is enough,
    # since only its raw bytes get copied, never validated.
    shell = tmp_path / swpw.VBA_SHELL_NAME
    openpyxl.Workbook().save(shell)
    with zipfile.ZipFile(shell, "a") as z:
        z.writestr("xl/vbaProject.bin", b"test-only, not a real vba project")
    return tmp_path


def hourly_trade(direction, counterparty, location, hours, mw, flow_date=FLOW):
    return {
        "direction": direction, "counterparty": counterparty, "location": location,
        "index": None, "price": 73.0, "is_monthly": False,
        "schedule": [(flow_date, h, float(mw)) for h in hours],
    }


BUY_LEG = hourly_trade("Buy", "AZPS", "PALOVERDE500", range(7, 23), 100)
SELL_LEG = hourly_trade("Sell", "BPAT", "MIDC", range(7, 23), 60)


def _run(trades=()):
    at = AppTest.from_file(PAGE_PATH, default_timeout=120)
    at.session_state["trades"] = list(trades)
    return at.run()


def _emit(at, event, seq=None):
    event = dict(event)
    event["seq"] = seq if seq is not None else at.session_state["mv_last_seq"] + 1
    at.session_state["mv_board"] = event
    return at.run()


def _click(at, label):
    [b for b in _in_dialog(at).button if b.label == label][0].click().run()
    return at

def _in_dialog(at):
    """Tell the page its modal is still open.

    A real click inside a dialog reruns the *fragment*; AppTest only does
    full script runs, and a full run with a dialog still flagged open is
    exactly what a dismissal looks like to
    ui.scheduling.state.dialog_was_dismissed. Without this, every step taken
    inside a modal here would read as "the trader pressed Esc".
    """
    at.session_state[BIDFILE_DIALOG] = False
    at.session_state[LINK_DIALOG] = False
    return at



def _link_to_swpw(at, frm="session:0", market="SWPW"):
    at = _emit(at, {"type": "link_request", "from": frm, "market": market})
    return _click(at, "Create link")


def _grid(at, event):
    """One bid-grid event, exactly as the component sends it — the code and
    the numbers all live in the grid now."""
    event = dict(event)
    event.setdefault("instance", "test-frame")
    event["seq"] = at.session_state["mv_bidgrid_last_seq"] + 1
    at.session_state["mv_bidgrid"] = event
    return _in_dialog(at).run()


def _fill_code(at, side, pse, code, line=0):
    """Type a GCA/LCA, leaving the price at the side's default."""
    return _grid(at, {"type": "code", "side": side, "pse": pse, "line": line, "value": code})


def _fill_line(at, side, pse, code, price, line=0):
    at = _fill_code(at, side, pse, code, line)
    return _grid(at, {"type": "price", "side": side, "pse": pse, "line": line, "value": price})


def _lines(at, side=SHORT, pse="AZPS", market="SWPW"):
    return at.session_state["mv_bidfile_splits"][(market, FLOW, side, pse)]


class TestOpeningThePopup:
    def test_clicking_the_swpw_chip_opens_it(self, no_bilateral_db):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert not at.exception, [e.value for e in at.exception]
        assert at.session_state["mv_bidfile_market"] == "SWPW"
        assert list(_lines(at)[0]["mw_by_hour"].values()) == [100.0] * 16

    def test_clicking_an_unsupported_market_says_so_instead_of_nothing(self, no_bilateral_db):
        at = _emit(
            _run([BUY_LEG]),
            {"type": "link_request", "from": "session:0", "market": "AESO"},
        )
        _click(at, "Create link")
        at = _emit(at, {"type": "chip_click", "market": "AESO"})
        assert not at.exception, [e.value for e in at.exception]
        assert any("AESO" in i.value and "yet" in i.value.lower() for i in at.info)

    def test_a_market_with_no_open_position_says_so(self, no_bilateral_db):
        at = _emit(_run([BUY_LEG]), {"type": "chip_click", "market": "SWPW"})
        assert not at.exception, [e.value for e in at.exception]
        assert any("No open position" in i.value for i in at.info)

    def test_a_click_never_starts_a_link(self, no_bilateral_db):
        # The click/drag threshold in the frontend is what makes this safe;
        # this just confirms the Python side agrees a bare click never
        # creates a link on its own.
        at = _emit(_run([BUY_LEG]), {"type": "chip_click", "market": "SWPW"})
        assert at.session_state["mv_links"] == []
        assert at.session_state["mv_pending"] is None

    def test_opening_the_bidfile_closes_any_link_popup(self, no_bilateral_db):
        second = hourly_trade("Sell", "BPAT", "MIDC", range(7, 23), 60)
        at = _emit(
            _run([BUY_LEG, second]),
            {"type": "link_request", "from": "session:0", "to": "session:1"},
        )
        assert at.session_state["mv_pending"] is not None
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert at.session_state["mv_pending"] is None


class TestDismissingIt:
    """st.dialog gives no callback when a modal is closed with x, Esc or a
    click outside, so the state saying "this is open" used to survive it —
    and the next page run put the dialog straight back. Reported as
    "clicking Refresh systematically opens the SWPW popup", and reproduced
    in a browser exactly that way."""

    def test_a_page_run_with_nothing_touched_closes_it(self, no_bilateral_db):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert at.session_state["mv_bidfile_market"] == "SWPW"

        at.run()  # what a dismissal leaves behind: a page run, nothing touched
        assert at.session_state["mv_bidfile_market"] is None

    def test_so_refreshing_afterwards_does_not_bring_it_back(self, no_bilateral_db):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at.run()                       # dismissed
        _click(at, "↻")              # Refresh
        assert at.session_state["mv_bidfile_market"] is None
        assert not at.exception, [e.value for e in at.exception]

    def test_working_inside_it_keeps_it_open(self, no_bilateral_db):
        # The other half of the guard: a rerun the dialog asked for itself
        # must not read as a dismissal.
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _fill_code(at, SHORT, "AZPS", "PALOVERDE")
        assert at.session_state["mv_bidfile_market"] == "SWPW"
        assert _lines(at)[0]["code"] == "PALOVERDE"


class TestGenerating:
    def test_filling_in_the_one_line_and_generating_writes_a_real_file(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _fill_line(at, SHORT, "AZPS", "AZPS", -1.0)
        _click(at, "Generate Bid File")

        assert not at.exception, [e.value for e in at.exception]
        assert any("Bid file written" in s.value for s in at.success)
        path = redirect_swpw / "test-output.xlsm"
        assert path.is_file()
        # The regression this guards against: a plain (non-macro-enabled)
        # workbook saved to a .xlsm name opens in Excel as "the file format
        # or file extension is not valid" — not just a warning, a hard
        # refusal. Confirmed by a real user hitting it.
        with zipfile.ZipFile(path) as z:
            assert "macroEnabled.main" in z.read("[Content_Types].xml").decode("utf-8")
        ws = openpyxl.load_workbook(path)["BIDS"]
        assert ws["D5"].value == "SPP-SHORT(AZPS)"
        # BUY_LEG flows PPT HE7-22; the sheet is EPT, so HE7 PPT sits on
        # EPT HE10 and HE22 PPT on the first starred row.
        assert ws.cell(row=swpw.ppt_hour_row(7), column=4).value == 100.0
        assert ws.cell(row=swpw.ppt_hour_row(7), column=5).value == -1.0
        assert ws.cell(row=swpw.ppt_hour_row(22), column=4).value == 100.0
        assert swpw.ppt_hour_row(22) == swpw.STAR_ROWS[0]
        # openpyxl always reads a date cell back as datetime.
        assert ws["C5"].value == datetime(FLOW.year, FLOW.month, FLOW.day)

    def test_untouched_prices_go_out_at_the_side_s_default(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG, SELL_LEG]))
        at = _link_to_swpw(at, frm="session:1")
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _fill_code(at, SHORT, "AZPS", "AZPS")
        at = _fill_code(at, LONG, "BPAT", "MIDC")
        _click(at, "Generate Bid File")

        assert not at.exception, [e.value for e in at.exception]
        ws = openpyxl.load_workbook(redirect_swpw / "test-output.xlsm")["BIDS"]
        # Shorts first from column D, then longs: MW/Price, MW/Price.
        assert ws["D5"].value == "SPP-SHORT(AZPS)"
        assert ws["F5"].value == "SPP-LONG(MIDC)"
        row = swpw.ppt_hour_row(7)
        assert ws.cell(row=row, column=5).value == 0.0   # a short opens at 0
        assert ws.cell(row=row, column=7).value == 50.0  # a long at 50

    def test_generating_without_a_code_blocks_and_writes_nothing(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        _click(at, "Generate Bid File")
        assert any("code" in e.value for e in at.error)
        assert not (redirect_swpw / "test-output.xlsm").exists()

    def test_a_second_generate_asks_before_overwriting(self, no_bilateral_db, redirect_swpw):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _fill_line(at, SHORT, "AZPS", "AZPS", -1.0)
        _click(at, "Generate Bid File")
        _click(at, "Generate Bid File")

        assert any("already exists" in w.value for w in at.warning)
        assert not [b for b in at.button if b.label == "Overwrite and Generate"]

        _in_dialog(at).checkbox(key="mv_bidfile_overwrite_confirm").set_value(True).run()
        assert [b for b in at.button if b.label == "Overwrite and Generate"]
        _click(at, "Overwrite and Generate")
        assert not at.exception, [e.value for e in at.exception]
        assert any("Bid file written" in s.value for s in at.success)


class TestWhatTheGridPutsInTheFile:
    """The grid's own behaviour is covered in test_bidgrid.py; this is the
    part that matters here — an edit made in it reaching the workbook."""

    def test_a_price_typed_into_the_grid_reaches_every_hour_it_flows(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _fill_line(at, SHORT, "AZPS", "AZPS", -1.0)
        _click(at, "Generate Bid File")

        assert not at.exception, [e.value for e in at.exception]
        ws = openpyxl.load_workbook(redirect_swpw / "test-output.xlsm")["BIDS"]
        # One price per bid line, so it lands on every hour the position
        # flows — and on none of the ones it doesn't.
        assert ws.cell(row=swpw.ppt_hour_row(7), column=5).value == -1.0
        assert ws.cell(row=swpw.ppt_hour_row(22), column=5).value == -1.0
        assert ws.cell(row=swpw.ppt_hour_row(1), column=5).value is None

    def test_a_split_typed_into_the_grid_becomes_two_columns(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _grid(at, {"type": "split", "side": SHORT, "pse": "AZPS", "line": 0})
        for hour in range(7, 23):
            at = _grid(at, {"type": "mw", "side": SHORT, "pse": "AZPS", "line": 1,
                            "hour": hour, "value": 40.0})
        at = _fill_code(at, SHORT, "AZPS", "AZPS", line=0)
        at = _fill_code(at, SHORT, "AZPS", "TEPC", line=1)
        assert not at.error  # it reconciles: 60 + 40 is the 100 that traded
        _click(at, "Generate Bid File")

        assert not at.exception, [e.value for e in at.exception]
        ws = openpyxl.load_workbook(redirect_swpw / "test-output.xlsm")["BIDS"]
        assert ws["D5"].value == "SPP-SHORT(AZPS)"
        assert ws["F5"].value == "SPP-SHORT(TEPC)"
        assert ws.cell(row=swpw.ppt_hour_row(7), column=4).value == 60.0
        assert ws.cell(row=swpw.ppt_hour_row(7), column=6).value == 40.0


class TestSplitting:
    def test_the_plus_on_a_code_adds_a_second_empty_line(self, no_bilateral_db):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        at = _grid(at, {"type": "split", "side": SHORT, "pse": "AZPS", "line": 0})
        lines = _lines(at)
        assert len(lines) == 2
        assert lines[1]["mw_by_hour"] == {}

    def test_a_reconciled_split_generates_two_columns(self, no_bilateral_db, redirect_swpw):
        at = _link_to_swpw(_run([BUY_LEG]))
        # Seeded directly rather than driven through the hour grid widget —
        # domain/bidfiles.py's own tests cover the reconciliation math
        # itself; this only checks the page wires that state through to a
        # real file correctly.
        at.session_state["mv_bidfile_splits"] = {
            ("SWPW", FLOW, "SHORT", "AZPS"): [
                {"code": "AZPS", "price": -1.0, "mw_by_hour": {h: 60.0 for h in range(7, 23)}},
                {"code": "TEPC", "price": 0.0, "mw_by_hour": {h: 40.0 for h in range(7, 23)}},
            ]
        }
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert not at.error  # the split already reconciles — no complaint
        _click(at, "Generate Bid File")

        assert not at.exception, [e.value for e in at.exception]
        ws = openpyxl.load_workbook(redirect_swpw / "test-output.xlsm")["BIDS"]
        assert ws["D5"].value == "SPP-SHORT(AZPS)"
        assert ws["F5"].value == "SPP-SHORT(TEPC)"

    def test_an_unreconciled_split_is_reported_and_blocks_generation(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at.session_state["mv_bidfile_splits"] = {
            ("SWPW", FLOW, "SHORT", "AZPS"): [
                {"code": "AZPS", "price": -1.0, "mw_by_hour": {7: 60.0}},
                {"code": "TEPC", "price": 0.0, "mw_by_hour": {7: 10.0}},  # 100 expected, only 70
            ]
        }
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert any("HE7" in e.value for e in at.error)
        _click(at, "Generate Bid File")
        assert not (redirect_swpw / "test-output.xlsm").exists()


class TestScopeAcrossFiltersAndHiding:
    def test_the_bidfile_sees_a_position_even_when_its_square_is_filtered_out(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at.multiselect(key="mv_pse").set_value(["NOBODY"]).run()
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert not at.exception, [e.value for e in at.exception]
        assert list(_lines(at)[0]["mw_by_hour"].values()) == [100.0] * 16

    def test_the_bidfile_sees_a_position_even_when_its_square_is_hidden(
        self, no_bilateral_db, redirect_swpw
    ):
        at = _link_to_swpw(_run([BUY_LEG]))
        at = _emit(at, {"type": "dismiss", "key": "session:0"})
        at = _emit(at, {"type": "chip_click", "market": "SWPW"})
        assert not at.exception, [e.value for e in at.exception]
        assert list(_lines(at)[0]["mw_by_hour"].values()) == [100.0] * 16
