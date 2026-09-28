"""data/tagfiles/bilateral.py — filling the desk's blank tag template.

Every cell reference asserted here was read off the desk's own files (all
2,638 of them put the same label in the same cell), so a test failing here
means the output stopped matching what a scheduler opens.

Every write goes to tmp_path via the `redirect_tags` fixture — never
Y:\\West.
"""

from datetime import date

import openpyxl
import pytest

import data.tagfiles.bilateral as tagfile
from domain.tags import blank_tag

FLOW = date(2026, 9, 25)


@pytest.fixture
def redirect_tags(tmp_path, monkeypatch):
    monkeypatch.setattr(tagfile, "TAG_FOLDER", tmp_path)
    return tmp_path


def a_tag(**over):
    tag = blank_tag(FLOW)
    tag["name"] = "ABEX-SWPW"
    tag["label"] = "SOURCE-SWPW(WAUW)"
    tag["source"].update(
        {"market": "ABEX", "gca": "GWA", "point": "RIMROCK", "pse": "RRWE01",
         "comment": "call first", "contract": "WSPP"}
    )
    tag["sink"].update(
        {"market": "SPP", "lca": "SWPW", "point": "SINKNODE", "pse": "MAG001",
         "comment": "", "contract": ""}
    )
    tag["market_path"] = [
        {"pse": "RRWE01", "product": "G-F", "contract": ""},
        {"pse": "ABEX", "product": "", "contract": ""},
        {"pse": "MAG001", "product": "L", "contract": ""},
    ]
    tag["transmissions"] = [
        {"path": "WS/NWMT/NWMT-WAUW/MATL.NWMT-CROSSOVER/", "reservation": "109267045", "mw": ""},
        {"path": "FCATBTEP", "reservation": "", "mw": "50"},
    ]
    tag["carbon_copy"] = [{"type": "BA", "market": "EPE"}]
    tag["mw_by_hour"] = {h: 25.0 for h in range(1, 25)}
    tag.update(over)
    return tag


def sheet(tag):
    return tagfile.build_workbook(tag)["BIDS"]


class TestTheHeaderAndTheTwoHalves:
    def test_the_flow_date_and_the_header_line(self):
        ws = sheet(a_tag())
        assert ws["C5"].value == FLOW
        assert ws["D5"].value == "SOURCE-SWPW(WAUW)"

    def test_the_date_keeps_the_templates_own_format(self):
        assert sheet(a_tag())["C5"].number_format == "mm-dd-yy"

    def test_the_source_half(self):
        ws = sheet(a_tag())
        assert [ws[c].value for c in ("D7", "D8", "D9", "D10", "D11", "D12")] == [
            "ABEX", "GWA", "RIMROCK", "RRWE01", "call first", "WSPP"
        ]

    def test_the_sink_half(self):
        ws = sheet(a_tag())
        assert [ws[c].value for c in ("D44", "D45", "D46", "D47")] == [
            "SPP", "SWPW", "SINKNODE", "MAG001"
        ]

    def test_a_blank_field_is_left_blank_not_written_as_empty(self):
        ws = sheet(a_tag())
        assert ws["D48"].value is None
        assert ws["D49"].value is None


class TestTheHourColumnIsEastern:
    """The sheet's hour column is EPT and the app is PPT — the same +3 the
    SWPW bid file needs, and the reason both have `1*`/`2*`/`3*` rows under
    HE24. Confirmed against the desk's files: "… TEPC HE7-14" fills rows
    23-30, which is EPT HE10-17."""

    def test_ppt_he1_lands_three_rows_down(self):
        ws = sheet(a_tag(mw_by_hour={1: 40.0}))
        assert ws.cell(17, 4).value == 40.0  # EPT HE4
        assert ws.cell(14, 4).value is None

    def test_the_desks_own_he7_to_14_position(self):
        ws = sheet(a_tag(mw_by_hour={h: 30.0 for h in range(7, 15)}))
        filled = [r for r in range(14, 41) if ws.cell(r, 4).value is not None]
        assert filled == list(range(23, 31))

    def test_the_last_three_hours_of_the_day_land_on_the_stars(self):
        ws = sheet(a_tag(mw_by_hour={22: 1.0, 23: 2.0, 24: 3.0}))
        assert [ws.cell(r, 4).value for r in (38, 39, 40)] == [1.0, 2.0, 3.0]

    def test_a_full_day_fills_exactly_twenty_four_slots(self):
        ws = sheet(a_tag())
        filled = [r for r in range(14, 41) if ws.cell(r, 4).value is not None]
        assert filled == list(range(17, 41))

    def test_both_halves_carry_the_same_schedule(self):
        ws = sheet(a_tag(mw_by_hour={7: 12.0, 24: 3.0}))
        assert ws.cell(23, 4).value == 12.0 and ws.cell(60, 4).value == 12.0
        assert ws.cell(40, 4).value == 3.0 and ws.cell(77, 4).value == 3.0

    def test_a_zero_hour_is_left_blank(self):
        # Matching every real file: an hour that doesn't flow is empty,
        # never an explicit 0.
        ws = sheet(a_tag(mw_by_hour={1: 0.0, 2: 5.0}))
        assert ws.cell(17, 4).value is None
        assert ws.cell(18, 4).value == 5.0


class TestTheThreeListBlocks:
    def test_transmissions_start_at_row_80(self):
        ws = sheet(a_tag())
        assert ws["D80"].value == "WS/NWMT/NWMT-WAUW/MATL.NWMT-CROSSOVER/"
        assert ws["D81"].value == "FCATBTEP"

    def test_a_reservation_number_is_written_as_the_number_it_is(self):
        # Every real file holds these as integers, not text.
        assert sheet(a_tag())["E80"].value == 109267045

    def test_a_contract_reference_in_that_column_stays_as_typed(self):
        tag = a_tag(transmissions=[{"path": "x", "reservation": "FCEX4EPE", "mw": ""}])
        assert sheet(tag)["E80"].value == "FCEX4EPE"

    def test_the_market_path_starts_at_row_94(self):
        ws = sheet(a_tag())
        assert [(ws.cell(r, 4).value, ws.cell(r, 5).value) for r in (94, 95, 96)] == [
            ("RRWE01", "G-F"), ("ABEX", None), ("MAG001", "L")
        ]

    def test_the_carbon_copy_starts_at_row_105(self):
        ws = sheet(a_tag())
        assert (ws["D105"].value, ws["E105"].value) == ("BA", "EPE")

    def test_an_empty_row_is_dropped_rather_than_leaving_a_gap(self):
        tag = a_tag(
            carbon_copy=[{"type": "", "market": ""}, {"type": "BA", "market": "EPE"}]
        )
        assert sheet(tag)["D105"].value == "BA"


class TestTheTemplateItself:
    def test_its_labels_survive(self):
        ws = sheet(a_tag())
        assert ws["C1"].value == "Bilateral Tag"
        assert (ws["C8"].value, ws["C45"].value) == ("GCA", "LCA")
        assert (ws["D79"].value, ws["C93"].value) == ("Path", "Market Path Product")
        assert ws["C104"].value == "Carbon Copy"

    def test_the_desks_own_cell_comments_come_along(self):
        ws = sheet(a_tag())
        assert ws["C39"].comment is not None

    def test_building_a_tag_never_writes_to_the_template(self):
        before = tagfile.TEMPLATE.read_bytes()
        tagfile.build_workbook(a_tag())
        assert tagfile.TEMPLATE.read_bytes() == before

    def test_two_tags_do_not_bleed_into_each_other(self):
        first = sheet(a_tag())
        second = sheet(blank_tag(FLOW) | {"mw_by_hour": {}, "name": "empty"})
        assert first["D7"].value == "ABEX"
        assert second["D7"].value is None


class TestWhereItLands:
    def test_while_under_review_everything_goes_to_the_test_folder(self, redirect_tags):
        assert tagfile.TEST_MODE, "flip this test's expectation when going live"
        assert tagfile.day_folder(FLOW) == redirect_tags / "Test"

    def test_live_it_picks_the_day_folder_the_desk_already_has(
        self, redirect_tags, monkeypatch
    ):
        monkeypatch.setattr(tagfile, "TEST_MODE", False)
        (redirect_tags / "25-26 September 2026").mkdir()
        (redirect_tags / "Templates").mkdir()
        assert tagfile.day_folder(FLOW) == redirect_tags / "25-26 September 2026"

    def test_an_exact_day_folder_wins_over_a_copy_of_it(
        self, redirect_tags, monkeypatch
    ):
        monkeypatch.setattr(tagfile, "TEST_MODE", False)
        (redirect_tags / "25 September 2026 - Copy").mkdir()
        (redirect_tags / "25 September 2026").mkdir()
        assert tagfile.day_folder(FLOW) == redirect_tags / "25 September 2026"

    def test_with_no_folder_for_the_day_it_names_a_new_one(
        self, redirect_tags, monkeypatch
    ):
        monkeypatch.setattr(tagfile, "TEST_MODE", False)
        assert tagfile.day_folder(FLOW) == redirect_tags / "25 September 2026"

    def test_an_unreachable_root_still_produces_a_path(self, monkeypatch, tmp_path):
        monkeypatch.setattr(tagfile, "TEST_MODE", False)
        monkeypatch.setattr(tagfile, "TAG_FOLDER", tmp_path / "not-mounted")
        assert tagfile.day_folder(FLOW).name == "25 September 2026"

    def test_the_full_target_path(self, redirect_tags):
        assert tagfile.target_path(a_tag()) == (
            redirect_tags / "Test"
            / "Tag_dynamique_bilateral - September 25 2026 - ABEX-SWPW.xlsx"
        )


class TestWriting:
    def test_it_writes_a_workbook_a_scheduler_can_open(self, redirect_tags):
        path = tagfile.write_tag_file(a_tag())
        assert path.exists()
        ws = openpyxl.load_workbook(path)["BIDS"]
        assert ws["D8"].value == "GWA"
        assert ws["D94"].value == "RRWE01"

    def test_it_creates_the_day_folder(self, redirect_tags):
        assert not (redirect_tags / "Test").exists()
        tagfile.write_tag_file(a_tag())
        assert (redirect_tags / "Test").is_dir()

    def test_it_refuses_to_replace_a_tag_that_is_already_there(self, redirect_tags):
        tagfile.write_tag_file(a_tag())
        with pytest.raises(FileExistsError):
            tagfile.write_tag_file(a_tag())

    def test_overwrite_is_something_the_caller_has_to_ask_for(self, redirect_tags):
        tagfile.write_tag_file(a_tag())
        path = tagfile.write_tag_file(a_tag(label="SECOND GO"), overwrite=True)
        assert openpyxl.load_workbook(path)["BIDS"]["D5"].value == "SECOND GO"

    def test_two_names_on_one_day_are_two_files(self, redirect_tags):
        tagfile.write_tag_file(a_tag())
        second = tagfile.write_tag_file(a_tag(name="ABEX-SWPW second leg"))
        assert second.name.endswith("ABEX-SWPW second leg.xlsx")
        assert len(list((redirect_tags / "Test").glob("*.xlsx"))) == 2
